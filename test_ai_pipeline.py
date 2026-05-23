import unittest
from io import BytesIO
import zipfile
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

import main
from backend.services.matcher import tfidf_cosine_score
from backend.services.resume_service import clean_text
from backend.services.skill_extractor import extract_skills


client = TestClient(main.app)


def make_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_table_bytes(text: str) -> bytes:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = text
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_header_bytes(header_text: str, body_text: str = "") -> bytes:
    document = Document()
    # python-docx creates an empty paragraph in a new header; write to it to avoid a leading blank line.
    section = document.sections[0]
    header = section.header
    if header.paragraphs:
        header.paragraphs[0].text = header_text
    else:
        header.add_paragraph(header_text)

    if body_text:
        document.add_paragraph(body_text)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_textbox_like_bytes(text: str) -> bytes:
    """
    Create a DOCX where text only exists inside a VML textbox structure inside mc:AlternateContent.

    `python-docx` usually ignores this content when traversing Document.paragraphs / tables, but our
    ZIP/XML fallback should still extract <w:t> nodes from document.xml.
    """
    base = Document()
    # Keep the main story empty (no visible paragraph text).
    output = BytesIO()
    base.save(output)
    base_bytes = output.getvalue()

    with zipfile.ZipFile(BytesIO(base_bytes), "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    doc_xml_bytes = parts.get("word/document.xml", b"")
    doc_xml = doc_xml_bytes.decode("utf-8", errors="ignore")

    injected = (
        "<w:p><w:r><mc:AlternateContent>"
        "<mc:Fallback>"
        "<w:pict><v:shape id=\"TextBox1\" type=\"#_x0000_t202\" style=\"width:200pt;height:50pt\">"
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>"
        + text +
        "</w:t></w:r></w:p></w:txbxContent></v:textbox>"
        "</v:shape></w:pict>"
        "</mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>"
    )

    # Insert before the section properties element at end-of-body.
    insert_at = doc_xml.rfind("<w:sectPr")
    if insert_at == -1:
        insert_at = doc_xml.rfind("</w:body>")
    if insert_at == -1:
        raise RuntimeError("Unexpected DOCX document.xml structure (missing w:sectPr/w:body end)")

    doc_xml = doc_xml[:insert_at] + injected + doc_xml[insert_at:]
    parts["word/document.xml"] = doc_xml.encode("utf-8")

    rebuilt = BytesIO()
    with zipfile.ZipFile(rebuilt, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, blob in parts.items():
            zf.writestr(name, blob)

    return rebuilt.getvalue()


class AiPipelineTests(unittest.TestCase):
    def setUp(self):
        self.email = f"pipeline-{uuid4().hex}@example.com"
        self.password = "Pass12345"
        client.post("/auth/register", json={"email": self.email, "password": self.password})
        login = client.post("/auth/login", json={"email": self.email, "password": self.password})
        self.headers = {"Authorization": f"Bearer {login.json()['data']['access_token']}"}

    def test_clean_text_normalizes_noise(self):
        cleaned = clean_text("  PYTHON!!!\nFastAPI\t\tSQL --- Docker  ")

        self.assertEqual(cleaned, "python fastapi sql docker")

    def test_skill_extraction_is_unique_and_accurate(self):
        skills = extract_skills("Python python FastAPI SQL Node.js sklearn")

        self.assertEqual(skills.count("python"), 1)
        self.assertIn("fastapi", skills)
        self.assertIn("sql", skills)
        self.assertIn("node js", skills)
        self.assertIn("scikit learn", skills)

    def test_tfidf_score_is_clamped_to_zero_to_100(self):
        score = tfidf_cosine_score("Python FastAPI SQL", "Python FastAPI SQL", ["python", "fastapi", "sql"])

        self.assertGreaterEqual(score, 0)
        self.assertLessEqual(score, 100)

    def test_parse_resume_supports_docx(self):
        content = make_docx_bytes("Python FastAPI SQL Docker backend engineer")

        response = client.post(
            "/parse-resume",
            headers=self.headers,
            files={
                "file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertIn("python", data["text"])
        self.assertIn("fastapi", data["skills"])
        profile = data.get("profile") or {}
        # Name extraction is best-effort; basic shape should always exist.
        self.assertIn("candidate_name", profile)
        self.assertIn("name_confidence", profile)

    def test_docx_upload_then_match_pipeline(self):
        candidate_id = f"docx-{uuid4().hex}"
        content = make_docx_bytes("Python FastAPI SQL Docker backend engineer")

        upload_response = client.post(
            "/upload-resume",
            headers=self.headers,
            data={"candidate_id": candidate_id},
            files={
                "file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(upload_response.status_code, 200)
        self.assertEqual(upload_response.json()["data"]["candidate_id"], candidate_id)

        match_response = client.post(
            "/match",
            headers=self.headers,
            json={"job_description": "Python FastAPI SQL Docker"},
        )

        self.assertEqual(match_response.status_code, 200)
        candidates = match_response.json()["data"]["candidates"]
        matched_ids = {candidate["candidate_id"] for candidate in candidates}
        self.assertIn(candidate_id, matched_ids)


    def test_parse_resume_supports_docx_tables(self):
        content = make_docx_table_bytes("Kubernetes Terraform AWS")

        response = client.post(
            "/parse-resume",
            headers=self.headers,
            files={
                "file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertIn("kubernetes", data["text"])


    def test_parse_resume_supports_docx_headers(self):
        content = make_docx_header_bytes("Jane Doe - Data Engineer", "Python Spark SQL")

        response = client.post(
            "/parse-resume",
            headers=self.headers,
            files={
                "file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertIn("jane", data["text"])
        self.assertIn("spark", data["text"])
        profile = data.get("profile") or {}
        # The header line includes a job title; extractor should avoid taking the whole string,
        # but should still find the human name portion when present.
        extracted_name = (profile.get("candidate_name") or "").lower()
        self.assertTrue(
            "jane" in extracted_name or extracted_name == "",
            msg=f"Unexpected extracted name: {extracted_name!r} profile={profile!r}",
        )


    def test_parse_resume_supports_docx_textbox_like_xml_via_fallback(self):
        content = make_docx_textbox_like_bytes("InjectedTextBoxContent")

        response = client.post(
            "/parse-resume",
            headers=self.headers,
            files={
                "file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertIn("injectedtextboxcontent".lower(), data["text"])
        # Ensure fallback ran (primary python-docx often returns empty for textbox-only docs).
        self.assertEqual(data.get("extraction", {}).get("parser_used", ""), "docx-xml")


if __name__ == "__main__":
    unittest.main()
