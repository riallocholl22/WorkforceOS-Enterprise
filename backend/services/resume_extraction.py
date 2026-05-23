import hashlib
import logging
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from typing import Any

try:
    from PyPDF2 import PdfReader
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    import pdfplumber
except Exception:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    from docx import Document
except Exception:  # pragma: no cover - optional dependency
    Document = None

try:
    import pytesseract  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    pytesseract = None

try:
    from PIL import Image  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    Image = None

try:
    import pypdfium2  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    pypdfium2 = None


logger = logging.getLogger(__name__)

MAX_DOCX_XML_BYTES = 6 * 1024 * 1024
MAX_DOCX_FILES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_OCR_PDF_PAGES = 3
MAX_OCR_DOCX_IMAGES = 4
MAX_OCR_TEXT_CHARS = 12000


def _sha256_bytes(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()


def normalize_whitespace(text: str) -> str:
    text = str(text or "")
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_for_matching(text: str) -> str:
    # Normalized form optimized for keyword/semantic matching (keeps +, #, ., @).
    text = normalize_whitespace(text).lower()
    text = re.sub(r"[^a-z0-9@+#.\s-]", " ", text)
    text = text.replace("-", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _confidence_score(text: str) -> float:
    text = normalize_whitespace(text)
    if not text:
        return 0.0
    length = len(text)
    alpha = sum(1 for ch in text if ch.isalpha())
    ratio = alpha / max(1, len(text))
    base = 0.25
    if length >= 250:
        base = 0.45
    if length >= 900:
        base = 0.65
    if length >= 2400:
        base = 0.8
    # Down-weight very low alphabetic density (common in broken extraction).
    if ratio < 0.22:
        base *= 0.6
    return max(0.05, min(0.95, round(base, 2)))


@dataclass
class ExtractionResult:
    text_raw: str = ""
    text_normalized: str = ""
    parser_used: str = ""
    file_hash: str = ""
    confidence: float = 0.0
    is_scanned_pdf: bool = False
    is_scanned_docx: bool = False
    ocr: dict[str, Any] = field(default_factory=dict)
    diagnostics: dict[str, Any] = field(default_factory=dict)

    def to_public_dict(self) -> dict[str, Any]:
        # Never return raw internals that could include stack traces.
        return {
            "parser_used": self.parser_used,
            "confidence": self.confidence,
            "is_scanned_pdf": self.is_scanned_pdf,
            "is_scanned_docx": self.is_scanned_docx,
            "ocr": self.ocr,
            "diagnostics": self.diagnostics,
            "text_length": len(self.text_raw or ""),
            "file_hash_prefix": (self.file_hash or "")[:16],
        }


def extract_text_from_txt(content: bytes) -> str:
    # Best-effort decode (handles mixed encodings without hard failing).
    for encoding in ("utf-8", "utf-16", "cp1252", "latin1"):
        try:
            decoded = content.decode(encoding)
            return normalize_whitespace(decoded)
        except Exception:
            continue
    return normalize_whitespace(content.decode("utf-8", errors="ignore"))


def _docx_container_text(container) -> tuple[list[str], dict[str, Any]]:
    paragraphs = getattr(container, "paragraphs", None) or []
    tables = getattr(container, "tables", None) or []
    paragraph_texts = [p.text for p in paragraphs if getattr(p, "text", "").strip()]
    cell_texts: list[str] = []
    cell_count = 0
    for table in tables:
        for row in getattr(table, "rows", []) or []:
            for cell in getattr(row, "cells", []) or []:
                cell_count += 1
                text = getattr(cell, "text", "") or ""
                if text.strip():
                    cell_texts.append(text)
    diagnostics = {
        "paragraphs": len(paragraphs),
        "paragraph_text_count": len(paragraph_texts),
        "tables": len(tables),
        "table_cells_seen": cell_count,
        "table_cell_text_count": len(cell_texts),
    }
    return paragraph_texts + cell_texts, diagnostics


def _extract_docx_python_docx(file_path: str) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"parser": "python-docx", "available": Document is not None}
    if Document is None:
        diagnostics["error"] = "not_installed"
        return "", diagnostics

    try:
        document = Document(file_path)
    except Exception as exc:
        diagnostics["error"] = f"open_failed:{type(exc).__name__}"
        return "", diagnostics

    body_texts, body_diag = _docx_container_text(document)
    diagnostics["body"] = body_diag

    header_texts: list[str] = []
    footer_texts: list[str] = []
    section_count = 0
    header_diag: list[dict[str, Any]] = []
    footer_diag: list[dict[str, Any]] = []

    for section in getattr(document, "sections", []) or []:
        section_count += 1
        for attr in ("header", "first_page_header", "even_page_header"):
            part = getattr(section, attr, None)
            if not part:
                continue
            part_texts, part_diag = _docx_container_text(part)
            header_texts.extend(part_texts)
            header_diag.append({"part": attr, **part_diag})
        for attr in ("footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, attr, None)
            if not part:
                continue
            part_texts, part_diag = _docx_container_text(part)
            footer_texts.extend(part_texts)
            footer_diag.append({"part": attr, **part_diag})

    diagnostics["sections"] = section_count
    diagnostics["headers"] = header_diag
    diagnostics["footers"] = footer_diag

    text = normalize_whitespace("\n".join(body_texts + header_texts + footer_texts))
    diagnostics["text_length"] = len(text)
    return text, diagnostics


def _docx_extract_wt_nodes(xml_bytes: bytes) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"parser": "docx-xml"}
    try:
        root = ET.fromstring(xml_bytes)
    except Exception as exc:
        diagnostics["error"] = f"xml_parse_failed:{type(exc).__name__}"
        return "", diagnostics

    texts: list[str] = []
    wt_count = 0
    for node in root.iter():
        tag = node.tag or ""
        local = tag.rsplit("}", 1)[-1] if "}" in tag else tag
        if local != "t":
            continue
        wt_count += 1
        if node.text:
            texts.append(node.text)
    diagnostics["w_t_nodes"] = wt_count
    combined = normalize_whitespace(" ".join(texts))
    return combined, diagnostics


def _safe_open_zip(file_path: str) -> tuple[zipfile.ZipFile | None, dict[str, Any]]:
    diagnostics: dict[str, Any] = {}
    try:
        zf = zipfile.ZipFile(file_path)
    except zipfile.BadZipFile:
        diagnostics["error"] = "bad_zip"
        return None, diagnostics
    except Exception as exc:
        diagnostics["error"] = f"zip_open_failed:{type(exc).__name__}"
        return None, diagnostics

    try:
        infos = zf.infolist()
        diagnostics["file_count"] = len(infos)
        if len(infos) > MAX_DOCX_FILES:
            diagnostics["error"] = "zip_too_many_files"
            zf.close()
            return None, diagnostics
        total_uncompressed = sum(int(info.file_size or 0) for info in infos)
        diagnostics["uncompressed_bytes"] = total_uncompressed
        if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
            diagnostics["error"] = "zip_uncompressed_too_large"
            zf.close()
            return None, diagnostics
    except Exception:
        # If we can't inspect safely, just proceed with best effort.
        diagnostics["warning"] = "zip_inspection_failed"

    return zf, diagnostics


def _extract_docx_zip_xml(file_path: str) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"parser": "zip+xml"}
    parts: list[str] = []
    part_diagnostics: list[dict[str, Any]] = []
    media_count = 0

    zf, zip_diag = _safe_open_zip(file_path)
    diagnostics["zip"] = zip_diag
    if zf is None:
        diagnostics["parts"] = []
        diagnostics["media_files"] = 0
        return "", diagnostics

    try:
        names = zf.namelist()
        media_count = sum(1 for name in names if name.startswith("word/media/"))
        candidates = []
        if "word/document.xml" in names:
            candidates.append("word/document.xml")
        candidates.extend(sorted(name for name in names if re.fullmatch(r"word/header\d+\.xml", name)))
        candidates.extend(sorted(name for name in names if re.fullmatch(r"word/footer\d+\.xml", name)))
        for optional in ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"):
            if optional in names:
                candidates.append(optional)

        # Bonus: scan additional word/*.xml parts (some templates store text in auxiliary story parts).
        extra_parts = [n for n in names if n.startswith("word/") and n.endswith(".xml") and n not in set(candidates)]
        candidates.extend(sorted(extra_parts)[:24])

        for part in candidates:
            try:
                info = zf.getinfo(part)
                if int(info.file_size or 0) > MAX_DOCX_XML_BYTES:
                    part_diagnostics.append({"part": part, "error": "part_too_large"})
                    continue
                xml_bytes = zf.read(part)
            except Exception as exc:
                part_diagnostics.append({"part": part, "error": f"read_failed:{type(exc).__name__}"})
                continue
            extracted, diag = _docx_extract_wt_nodes(xml_bytes)
            parts.append(extracted)
            part_diagnostics.append({"part": part, "text_length": len(extracted), **diag})
    finally:
        try:
            zf.close()
        except Exception:
            pass

    combined = normalize_whitespace("\n".join(part for part in parts if part))
    diagnostics["parts"] = part_diagnostics
    diagnostics["media_files"] = media_count
    diagnostics["text_length"] = len(combined)
    return combined, diagnostics


def extract_text_from_docx(file_path: str) -> tuple[str, dict[str, Any]]:
    # Primary: python-docx (body + tables + headers/footers)
    # Fallback: direct ZIP+XML <w:t> extraction (handles many textboxes/shapes)
    text1, diag1 = _extract_docx_python_docx(file_path)
    if text1.strip():
        return text1, {"selected": "python-docx", "primary": diag1}

    text2, diag2 = _extract_docx_zip_xml(file_path)
    if text2.strip():
        return text2, {"selected": "zip+xml", "primary": diag1, "fallback": diag2}

    return "", {"selected": "empty", "primary": diag1, "fallback": diag2}


def _extract_pdf_pypdf2(file_path: str) -> tuple[str, str]:
    if PdfReader is None:
        return "", "PyPDF2:not_installed"
    try:
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        text = normalize_whitespace("\n".join(pages))
        return text, "PyPDF2"
    except Exception as exc:
        return "", f"PyPDF2:error:{type(exc).__name__}"


def _extract_pdf_pdfplumber(file_path: str) -> tuple[str, dict[str, Any]]:
    if pdfplumber is None:
        return "", {"parser": "pdfplumber:not_installed"}
    diagnostics: dict[str, Any] = {"parser": "pdfplumber"}
    try:
        with pdfplumber.open(file_path) as pdf:
            page_texts: list[str] = []
            image_pages = 0
            textless_pages = 0
            for page in pdf.pages[:12]:
                page_text = page.extract_text() or ""
                if not page_text.strip():
                    textless_pages += 1
                page_texts.append(page_text)
                if getattr(page, "images", None):
                    if len(page.images) > 0:
                        image_pages += 1
            diagnostics["pages_checked"] = min(12, len(pdf.pages))
            diagnostics["image_pages"] = image_pages
            diagnostics["textless_pages"] = textless_pages
            text = normalize_whitespace("\n".join(page_texts))
            return text, diagnostics
    except Exception as exc:
        diagnostics["error"] = f"{type(exc).__name__}"
        return "", diagnostics


def _extract_pdf_pdfium_text(file_path: str) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"parser": "pypdfium2", "available": pypdfium2 is not None}
    if pypdfium2 is None:
        return "", diagnostics

    try:
        pdf = pypdfium2.PdfDocument(file_path)
        page_count = len(pdf)
        diagnostics["pages"] = page_count
        page_texts: list[str] = []
        for i in range(min(page_count, 8)):
            page = pdf[i]
            textpage = page.get_textpage()
            text_all = textpage.get_text_range()
            page_texts.append(text_all or "")
            try:
                textpage.close()
            except Exception:
                pass
            try:
                page.close()
            except Exception:
                pass
        try:
            pdf.close()
        except Exception:
            pass
        text = normalize_whitespace("\n".join(page_texts))
        diagnostics["text_length"] = len(text)
        return text, diagnostics
    except Exception as exc:
        diagnostics["error"] = f"{type(exc).__name__}"
        return "", diagnostics


def _render_pdf_pages_to_images(file_path: str, max_pages: int = MAX_OCR_PDF_PAGES) -> tuple[list[Any], dict[str, Any]]:
    diagnostics: dict[str, Any] = {"engine": "pypdfium2", "available": pypdfium2 is not None}
    if pypdfium2 is None:
        return [], diagnostics
    if Image is None:
        diagnostics["error"] = "pillow_not_available"
        return [], diagnostics

    images: list[Any] = []
    try:
        pdf = pypdfium2.PdfDocument(file_path)
        page_count = len(pdf)
        diagnostics["pages"] = page_count
        for i in range(min(page_count, max_pages)):
            page = pdf[i]
            # Render at a reasonable scale (balance quality vs speed).
            bitmap = page.render(scale=2.0)
            pil = bitmap.to_pil()
            images.append(pil)
            try:
                bitmap.close()
            except Exception:
                pass
            try:
                page.close()
            except Exception:
                pass
        try:
            pdf.close()
        except Exception:
            pass
    except Exception as exc:
        diagnostics["error"] = f"render_failed:{type(exc).__name__}"
        return [], diagnostics

    diagnostics["images"] = len(images)
    return images, diagnostics


def _ocr_images(images: list[Any]) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"engine": "pytesseract", "available": pytesseract is not None}
    if pytesseract is None:
        diagnostics["error"] = "pytesseract_not_available"
        return "", diagnostics
    if not images:
        diagnostics["error"] = "no_images"
        return "", diagnostics

    texts: list[str] = []
    try:
        for img in images[:MAX_OCR_PDF_PAGES]:
            try:
                text = pytesseract.image_to_string(img)
            except Exception as exc:
                diagnostics.setdefault("errors", []).append(type(exc).__name__)
                text = ""
            if text and len("".join(texts)) < MAX_OCR_TEXT_CHARS:
                texts.append(text)
            try:
                # Free resources eagerly for large images.
                img.close()
            except Exception:
                pass
    except Exception as exc:
        diagnostics["error"] = f"ocr_failed:{type(exc).__name__}"
        return "", diagnostics

    combined = normalize_whitespace("\n".join(texts))
    diagnostics["text_length"] = len(combined)
    return combined, diagnostics


def _extract_docx_ocr(file_path: str) -> tuple[str, dict[str, Any]]:
    diagnostics: dict[str, Any] = {"engine": "docx_media+pytesseract", "available": pytesseract is not None and Image is not None}
    if pytesseract is None or Image is None:
        diagnostics["error"] = "ocr_not_available"
        return "", diagnostics

    import io

    zf, zip_diag = _safe_open_zip(file_path)
    diagnostics["zip"] = zip_diag
    if zf is None:
        diagnostics["error"] = zip_diag.get("error") or "zip_open_failed"
        return "", diagnostics

    images: list[Any] = []
    seen = 0
    try:
        for name in zf.namelist():
            if not name.startswith("word/media/"):
                continue
            if seen >= MAX_OCR_DOCX_IMAGES:
                break
            try:
                blob = zf.read(name)
            except Exception:
                continue
            seen += 1
            try:
                img = Image.open(io.BytesIO(blob))  # type: ignore[name-defined]
                images.append(img)
            except Exception:
                continue
    finally:
        try:
            zf.close()
        except Exception:
            pass

    diagnostics["images_seen"] = seen
    text, ocr_diag = _ocr_images(images)
    diagnostics["ocr"] = ocr_diag
    return text, diagnostics


def extract_resume(file_path: str, extension: str, content: bytes) -> ExtractionResult:
    result = ExtractionResult(file_hash=_sha256_bytes(content))
    size_bytes = len(content or b"")
    result.diagnostics.update({"extension": extension, "size_bytes": size_bytes})

    raw_text = ""
    chain: list[dict[str, Any]] = []

    if extension == ".txt":
        raw_text = extract_text_from_txt(content)
        result.parser_used = "text/plain"
        chain.append({"parser": "text/plain", "text_length": len(raw_text)})

    elif extension == ".docx":
        raw_text, docx_diag = extract_text_from_docx(file_path)
        result.diagnostics["docx"] = docx_diag
        selected = docx_diag.get("selected")
        if selected == "python-docx":
            result.parser_used = "python-docx"
        elif selected == "zip+xml":
            result.parser_used = "docx-xml"
        else:
            result.parser_used = "docx_fallback_chain_empty"

        primary_len = (docx_diag.get("primary") or {}).get("text_length", 0) if isinstance(docx_diag.get("primary"), dict) else 0
        fallback_len = (docx_diag.get("fallback") or {}).get("text_length", 0) if isinstance(docx_diag.get("fallback"), dict) else 0
        chain.append({"parser": "python-docx", "text_length": int(primary_len or 0)})
        chain.append({"parser": "docx-xml", "text_length": int(fallback_len or 0)})

        media_files = (docx_diag.get("fallback") or {}).get("media_files") if isinstance(docx_diag.get("fallback"), dict) else 0
        if not raw_text.strip() and int(media_files or 0) > 0:
            result.is_scanned_docx = True
            # Best-effort OCR if available.
            if pytesseract is not None and Image is not None:
                try:
                    ocr_text, ocr_diag = _extract_docx_ocr(file_path)
                    result.diagnostics["docx_ocr"] = ocr_diag
                    if ocr_text.strip():
                        raw_text = ocr_text
                        result.parser_used = "docx-ocr"
                        result.ocr = {"status": "used", "engine": "pytesseract", "text_length": len(ocr_text)}
                    else:
                        result.ocr = {"status": "needed", "message": "DOCX appears image-based. OCR did not yield extractable text."}
                except Exception:
                    result.ocr = {"status": "needed", "message": "DOCX appears image-based. OCR is recommended."}
            else:
                result.ocr = {
                    "status": "needed",
                    "message": "DOCX appears image-based. OCR support is required for full extraction.",
                }

    elif extension == ".pdf":
        text1, status1 = _extract_pdf_pypdf2(file_path)
        chain.append({"parser": status1, "text_length": len(text1)})
        raw_text = text1

        if not raw_text.strip():
            text2, diag2 = _extract_pdf_pdfplumber(file_path)
            chain.append({"parser": diag2.get("parser"), "text_length": len(text2), **{k: v for k, v in diag2.items() if k != "parser"}})
            raw_text = text2
            if not raw_text.strip():
                # Fallback: pdfium text extraction (handles some complex layouts better than PyPDF2).
                text3, diag3 = _extract_pdf_pdfium_text(file_path)
                result.diagnostics["pdfium"] = diag3
                chain.append({"parser": "pypdfium2", "text_length": len(text3), **{k: v for k, v in diag3.items() if k != "parser"}})
                raw_text = text3

            if not raw_text.strip():
                # Heuristic scanned detection: lots of images, very low/no text.
                image_pages = int(diag2.get("image_pages") or 0)
                pages_checked = int(diag2.get("pages_checked") or 0)
                if pages_checked and image_pages >= max(1, pages_checked // 2) and size_bytes > 80_000:
                    result.is_scanned_pdf = True
                    # OCR if available.
                    if pytesseract is not None and pypdfium2 is not None and Image is not None:
                        images, render_diag = _render_pdf_pages_to_images(file_path)
                        ocr_text, ocr_diag = _ocr_images(images)
                        result.diagnostics["pdf_render"] = render_diag
                        result.diagnostics["pdf_ocr"] = ocr_diag
                        chain.append({"parser": "ocr", "text_length": len(ocr_text)})
                        if ocr_text.strip():
                            raw_text = ocr_text
                            result.ocr = {"status": "used", "engine": "pytesseract", "text_length": len(ocr_text)}
                        else:
                            result.ocr = {"status": "needed", "message": "PDF appears image-based. OCR is required for full extraction."}
                    else:
                        result.ocr = {
                            "status": "needed",
                            "message": "This CV appears to be image-based or scanned. OCR support is required for full extraction.",
                            "engine": "pytesseract (optional)",
                        }
                else:
                    result.ocr = {"status": "not_triggered"}

        if text1.strip():
            result.parser_used = "PyPDF2"
        elif raw_text.strip() and (result.ocr or {}).get("status") == "used":
            result.parser_used = "pdf-ocr"
        elif raw_text.strip() and result.diagnostics.get("pdfium"):
            result.parser_used = "pypdfium2"
        elif raw_text.strip():
            result.parser_used = "pdfplumber"
        else:
            result.parser_used = "pdf_fallback_chain"
    else:
        result.parser_used = "unsupported"

    result.text_raw = normalize_whitespace(raw_text)
    result.text_normalized = normalize_for_matching(result.text_raw)
    result.confidence = _confidence_score(result.text_raw)
    result.diagnostics["chain"] = chain

    # Minimal structured log for ops debugging (no raw text).
    logger.info(
        "resume_extraction extension=%s parser=%s confidence=%s scanned_pdf=%s scanned_docx=%s text_len=%s size_bytes=%s hash_prefix=%s",
        extension,
        result.parser_used,
        result.confidence,
        result.is_scanned_pdf,
        result.is_scanned_docx,
        len(result.text_raw or ""),
        size_bytes,
        (result.file_hash or "")[:16],
    )

    return result


def extractor_health() -> dict[str, Any]:
    return {
        "pdf": {
            "PyPDF2": PdfReader is not None,
            "pdfplumber": pdfplumber is not None,
            "pypdfium2": pypdfium2 is not None,
            "ocr_ready": bool(pypdfium2 is not None and pytesseract is not None and Image is not None),
        },
        "docx": {"python-docx": Document is not None, "zip_xml_fallback": True},
        "txt": {"text": True},
        "ocr": {
            "pytesseract": pytesseract is not None,
            "pillow": Image is not None,
            "pypdfium2": pypdfium2 is not None,
        },
    }
